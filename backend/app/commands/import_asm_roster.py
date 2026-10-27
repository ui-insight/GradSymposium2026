"""Load a source response workbook as flagged, unpublished draft candidates."""

import argparse
import asyncio
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from app.commands.setup_asm import setup
from app.db.engine import async_session_factory, engine
from app.services.asm_roster import prepare_submissions, stage_submissions


async def import_workbook(path: Path, report_path: Path) -> dict:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Form Responses 1"]
        values = list(sheet.iter_rows(values_only=True))
        submissions = prepare_submissions(values, path.name, digest, sheet.title)
    finally:
        workbook.close()
    if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
        raise ValueError(
            "The workbook changed during import; retry the source snapshot"
        )
    event_id = await setup(activate=True)
    async with async_session_factory() as db:
        try:
            report = await stage_submissions(db, event_id, submissions)
            report_path.parent.mkdir(parents=True, exist_ok=True)
            # Write before commit so a report failure leaves the roster untouched.
            temporary = report_path.with_suffix(".pending")
            temporary.write_text(json.dumps(report, indent=2, ensure_ascii=False))
            await db.commit()
            temporary.replace(report_path)
            return report
        except Exception:
            await db.rollback()
            raise
        finally:
            await engine.dispose()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    report = asyncio.run(import_workbook(args.source.resolve(), args.report.resolve()))
    print(
        json.dumps(
            {key: value for key, value in report.items() if key != "records"}, indent=2
        )
    )


if __name__ == "__main__":
    main()
