"""Exercise roster issuance/publication through the real API on isolated SQLite."""

import asyncio
import datetime

import httpx
import pytest
import pytest_asyncio
from fastapi import FastAPI
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

import app.models  # noqa: F401
from app.api.deps import get_current_admin, get_db
from app.api.v1.router import api_router
from app.db.base import Base
from app.db.schema import ensure_runtime_schema
from app.models.event import Event
from app.models.project import Project
from app.models.project_number_reservation import ProjectNumberReservation
from app.models.user import User


@pytest_asyncio.fixture
async def context(tmp_path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'test.db'}")
    async with engine.begin() as connection:
        await connection.execute(text("PRAGMA foreign_keys=ON"))
        await connection.run_sync(Base.metadata.create_all)
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with sessions() as db:
        db.add(
            Event(
                Event_ID="asm",
                Event_Name="ASM 2026",
                Event_Date=datetime.date(2026, 10, 31),
                Is_Active=True,
            )
        )
        await db.commit()
    app = FastAPI()
    app.include_router(api_router)

    async def test_db():
        async with sessions() as db:
            yield db

    app.dependency_overrides[get_db] = test_db
    app.dependency_overrides[get_current_admin] = lambda: User(
        User_ID="test-admin", Username="test", Hashed_Password="unused"
    )
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        yield client, sessions, app
    await engine.dispose()


def poster(**changes):
    return {
        "Event_ID": "asm",
        "Project_Title": "Microbial communities",
        "Presenter_First_Name": "Ada",
        "Presenter_Last_Name": "Rivera",
        "Category": "Poster",
        **changes,
    }


async def create(client, **changes):
    response = await client.post("/api/v1/projects/", json=poster(**changes))
    assert response.status_code == 201, response.text
    return response.json()


async def upload(client, content, preview=False):
    response = await client.post(
        f"/api/v1/projects/import-csv?event_id=asm&dry_run={str(preview).lower()}",
        files={"file": ("roster.csv", content, "text/csv")},
    )
    assert response.status_code == 200, response.text
    return response.json()


async def publish(client, value=True):
    return await client.post(
        "/api/v1/projects/publication", json={"Event_ID": "asm", "Is_Published": value}
    )


async def directory(client):
    response = await client.get("/api/v1/public/posters")
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    return response.json()


@pytest.mark.asyncio
async def test_preview_rolls_back_and_repeat_import_is_idempotent(context):
    client, sessions, _ = context
    csv = (
        "Title,First_Name,Last_Name,Category\n"
        "Soil microbes,Ada,Rivera,Poster\nYeast,Alex,Lee,Poster\n"
    )
    preview = await upload(client, csv, preview=True)
    assert [item["Project_Number"] for item in preview["preview"]] == ["P-001", "P-002"]
    async with sessions() as db:
        assert await db.scalar(select(func.count()).select_from(Project)) == 0
        assert (
            await db.scalar(select(func.count()).select_from(ProjectNumberReservation))
            == 0
        )
    imported = await upload(client, csv)
    assert imported["imported"] == 2
    assert all(not item["Is_Published"] for item in imported["preview"])
    repeated = await upload(client, csv)
    assert repeated["errors"] == []
    assert (repeated["imported"], repeated["skipped"]) == (0, 2)
    assert (await directory(client))["posters"] == []


@pytest.mark.asyncio
async def test_invalid_or_duplicate_batch_is_atomic(context):
    client, sessions, _ = context
    for csv in (
        '"unterminated header',
        "Title,First_Name,Last_Name,Category\nValid,Ada,Lee,Poster\nInvalid,Bob,Lee,Talk\n",
        "Title,First_Name,Last_Name,Category\nSame,Ada,Lee,Poster\nSame,Ada,Lee,Poster\n",
        "Title,First_Name,Last_Name,Category\nMalformed,Ada,Lee,Poster,extra\n",
        "Title,Title,First_Name,Last_Name,Category\nFirst,Second,Ada,Lee,Poster\n",
    ):
        result = await upload(client, csv)
        assert result["errors"]
        assert result["imported"] == 0
    async with sessions() as db:
        assert await db.scalar(select(func.count()).select_from(Project)) == 0
        assert (
            await db.scalar(select(func.count()).select_from(ProjectNumberReservation))
            == 0
        )


@pytest.mark.asyncio
async def test_explicit_numbers_reserved_before_automatic_and_copresenters_preserved(
    context,
):
    client, _, _ = context
    csv = (
        "Title,Presenter_Display_Name,Category,Project_Number,Institution\n"
        'Soil,"Alex Chen, Morgan Rivera",Poster,,University of Idaho\n'
        "Yeast,Alex Lee,Poster,P-001,WSU\n"
    )
    result = await upload(client, csv)
    assert result["errors"] == []
    assert [item["Project_Number"] for item in result["preview"]] == ["P-002", "P-001"]
    assert (await publish(client)).status_code == 200
    posters = (await directory(client))["posters"]
    assert posters[1]["Presenter_Display_Name"] == "Alex Chen, Morgan Rivera"
    assert posters[1]["Institution"] == "University of Idaho"


@pytest.mark.asyncio
async def test_legacy_numbers_withdrawals_and_deletions_remain_reserved(context):
    client, sessions, _ = context
    async with sessions() as db:
        db.add(
            Project(
                **poster(Project_Number="P-99"), Project_ID="legacy", Is_Active=False
            )
        )
        await db.commit()
    first = await create(client)
    assert first["Project_Number"] == "P-100"
    assert (
        await client.delete(f"/api/v1/projects/{first['Project_ID']}")
    ).status_code == 204
    next_project = await create(client, Project_Title="Late addition")
    assert next_project["Project_Number"] == "P-101"
    response = await client.post(
        "/api/v1/projects/", json=poster(Project_Number="P-100")
    )
    assert response.status_code == 409
    assert (
        await client.patch(
            f"/api/v1/projects/{next_project['Project_ID']}",
            json={"Project_Number": "P-001"},
        )
    ).status_code == 409


@pytest.mark.asyncio
async def test_new_and_corrected_posters_remain_drafts_until_published(context):
    client, _, _ = context
    first = await create(client)
    assert (await directory(client))["event"] is None
    assert (await publish(client)).status_code == 200
    assert len((await directory(client))["posters"]) == 1
    await create(client, Project_Title="Late arrival")
    assert len((await directory(client))["posters"]) == 1
    project_url = f"/api/v1/projects/{first['Project_ID']}"
    response = await client.patch(
        project_url, json={"Institution": "Corrected institution"}
    )
    assert response.status_code == 200
    assert not response.json()["Is_Published"]
    assert (await directory(client))["posters"] == []
    await publish(client)
    await client.post(f"{project_url}/withdraw")
    assert len((await directory(client))["posters"]) == 1
    await client.post(f"{project_url}/reinstate")
    assert len((await directory(client))["posters"]) == 1
    await publish(client, False)
    assert (await directory(client))["event"] is None


@pytest.mark.asyncio
async def test_public_schema_excludes_private_data_and_other_events(context):
    client, sessions, _ = context
    await create(
        client,
        Presenter_Email="private@example.com",
        Advisor_Name="Private advisor",
        Institution="Idaho",
        Table_Number="B-17",
    )
    await create(client, Category="Art", Project_Title="Not a poster")
    await publish(client)
    async with sessions() as db:
        db.add(
            Event(
                Event_ID="old",
                Event_Name="Old GPSA",
                Event_Date=datetime.date(2026, 3, 13),
                Is_Active=False,
                Public_Poster_Lookup_Enabled=True,
            )
        )
        await db.flush()
        db.add(
            Project(
                **poster(
                    Event_ID="old", Project_Number="P-001", Project_Title="Other event"
                ),
                Is_Published=True,
            )
        )
        await db.commit()
    payload = await directory(client)
    assert len(payload["posters"]) == 1
    assert set(payload["posters"][0]) == {
        "Project_Number",
        "Project_Title",
        "Presenter_Display_Name",
        "Institution",
        "Table_Number",
    }
    assert "private@example.com" not in str(payload)
    assert "Private advisor" not in str(payload)
    assert "Other event" not in str(payload)
    assert set(payload["event"]) == {
        "Event_Name",
        "Event_Date",
        "Event_End_Date",
        "Location",
    }


@pytest.mark.asyncio
async def test_ambiguous_active_events_fail_closed(context):
    client, sessions, _ = context
    await create(client)
    await publish(client)
    async with sessions() as db:
        db.add(
            Event(
                Event_ID="other",
                Event_Name="Other",
                Event_Date=datetime.date(2026, 10, 31),
                Is_Active=True,
            )
        )
        await db.commit()
    assert (await directory(client))["event"] is None
    assert (await publish(client)).status_code == 409


@pytest.mark.asyncio
async def test_public_access_needs_no_auth_but_admin_mutations_do(context):
    client, _, app = context
    app.dependency_overrides.pop(get_current_admin)
    assert (await client.get("/api/v1/public/posters")).status_code == 200
    assert (await client.post("/api/v1/projects/", json=poster())).status_code == 401
    assert (await publish(client)).status_code == 401


@pytest.mark.asyncio
async def test_concurrent_number_issuance(context):
    client, _, _ = context
    projects = await asyncio.gather(
        *(create(client, Project_Title=f"Poster {index}") for index in range(8))
    )
    assert {project["Project_Number"] for project in projects} == {
        f"P-{index:03d}" for index in range(1, 9)
    }


@pytest.mark.asyncio
async def test_title_revisions_require_review_and_numeric_aliases_are_reserved(context):
    client, _, _ = context
    await create(client, Project_Number="P-01", Presenter_Email="ada@example.edu")
    result = await upload(
        client,
        "Title,First_Name,Last_Name,Category,Email\n"
        "A revised title,Ada,Rivera,Poster,ada@example.edu\n",
    )
    assert result["imported"] == 0
    assert "Possible title revision" in result["errors"][0]
    response = await client.post(
        "/api/v1/projects/", json=poster(Project_Number="P-001")
    )
    assert response.status_code == 409


def test_alembic_fresh_install_and_existing_schema_upgrade(tmp_path, monkeypatch):
    from pathlib import Path

    import sqlalchemy as sa
    from alembic import command
    from alembic.config import Config

    from app.config import settings

    path = tmp_path / "migration.db"
    monkeypatch.setattr(settings, "database_url", f"sqlite+aiosqlite:///{path}")
    config = Config(str(Path(__file__).parents[1] / "alembic.ini"))
    command.upgrade(config, "head")
    engine = sa.create_engine(f"sqlite:///{path}")
    with engine.begin() as connection:
        connection.execute(
            text(
                'INSERT INTO "Event" ("Event_ID", "Event_Name", "Event_Date", '
                '"Is_Active", "Created_At") '
                "VALUES ('preserved', 'Historical event', '2026-03-13', "
                "0, CURRENT_TIMESTAMP)"
            )
        )
    command.downgrade(config, "base")
    command.upgrade(config, "head")
    with engine.connect() as connection:
        assert (
            connection.scalar(text('SELECT "Event_Name" FROM "Event"'))
            == "Historical event"
        )
        assert not connection.scalar(
            text('SELECT "Public_Poster_Lookup_Enabled" FROM "Event"')
        )
        assert sa.inspect(connection).has_table("Project_Number_Reservation")
    engine.dispose()


@pytest.mark.asyncio
async def test_runtime_upgrade_is_additive_and_repeatable(tmp_path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'old.db'}")
    async with engine.begin() as connection:
        await connection.execute(
            text('CREATE TABLE "Event" ("Event_ID" TEXT PRIMARY KEY)')
        )
        await connection.execute(
            text(
                'CREATE TABLE "Project" ("Project_ID" TEXT PRIMARY KEY, '
                '"Event_ID" TEXT)'
            )
        )
        await connection.execute(
            text('CREATE TABLE "Score" ("Score_ID" INTEGER PRIMARY KEY)')
        )
        await connection.execute(
            text('CREATE TABLE "Feedback" ("Feedback_ID" INTEGER PRIMARY KEY)')
        )
        await connection.execute(
            text('INSERT INTO "Project" ("Project_ID") VALUES (\'historical\')')
        )
        await ensure_runtime_schema(connection)
        await ensure_runtime_schema(connection)
        result = (
            await connection.execute(
                text('SELECT "Project_ID", "Is_Published" FROM "Project"')
            )
        ).one()
        assert tuple(result) == ("historical", 0)
    await engine.dispose()


def source_fixture():
    return [
        [
            "Timestamp",
            "Abstract title ",
            "Presenting author",
            "Presenting author email address",
            "Abstract (250 word limit)",
            "Your current position",
            "Please check the appropriate box(es) if you would like "
            "to be considered for:",
            "Please indicate your subfield(s)/specialty:",
            "PI/ Lab name",
            "Institution",
        ],
        [
            1,
            "Repeated research",
            "Alex Lee",
            "alex@example.edu",
            "Original",
            "Graduate",
            "A poster at the main conference",
            "Microbiology",
            "Lab",
            "Idaho",
        ],
        [
            2,
            "Repeated research",
            "Alex Lee",
            "alex@example.edu",
            "Revision",
            "Graduate",
            "A poster at the main conference",
            "Microbiology",
            "Lab",
            "Idaho",
        ],
        [
            3,
            "Talk research",
            "Sam Gray",
            "sam@example.edu",
            "Abstract",
            "Faculty",
            "A talk at the main conference",
            "Microbiology",
            "Lab",
            "WSU",
        ],
        [
            4,
            "Joint research",
            "Jo Chen, Pat Rivera",
            "jo@example.edu, pat@example.edu",
            "Abstract",
            "Graduate",
            "A poster at the main conference",
            "Microbiology",
            "Lab",
            "Idaho",
        ],
    ]


def test_source_flags_retain_every_row_without_merging():
    from app.services.asm_roster import prepare_submissions

    submissions = prepare_submissions(
        source_fixture(), "source.xlsx", "hash", "Form Responses 1"
    )
    assert len(submissions) == 4
    assert [item["review"]["source_row"] for item in submissions] == [2, 3, 4, 5]
    assert submissions[0]["review"]["duplicate_rows"] == [2, 3]
    assert submissions[1]["review"]["duplicate_rows"] == [2, 3]
    assert submissions[2]["review"]["flags"] == ["talk_only"]
    assert submissions[3]["review"]["flags"] == ["co_presenters"]
    assert len({item["source_key"] for item in submissions}) == 4
    # Multiple presenters may share a single corresponding-author email.
    for name in ("Jo Chen, Pat Rivera", "Jo Chen and Pat Rivera"):
        values = source_fixture()
        values[-1][2] = name
        values[-1][3] = "jo@example.edu"
        prepared = prepare_submissions(
            values, "source.xlsx", "hash", "Form Responses 1"
        )
        assert "co_presenters" in prepared[-1]["review"]["flags"]


@pytest.mark.asyncio
async def test_staged_source_is_repeatable_and_publication_is_blocked(context):
    from app.services.asm_roster import prepare_submissions, stage_submissions

    client, sessions, _ = context
    submissions = prepare_submissions(
        source_fixture(), "source.xlsx", "hash", "Form Responses 1"
    )
    async with sessions() as db:
        first = await stage_submissions(db, "asm", submissions)
        await db.commit()
    assert first["source_submissions"] == 4
    assert first["review_required"] == 3
    assert first["duplicate_groups"] == [[2, 3]]
    assert [record["number"] for record in first["records"]] == [
        "P-001",
        "P-002",
        "P-003",
        "P-004",
    ]
    async with sessions() as db:
        repeated = await stage_submissions(db, "asm", submissions)
        await db.commit()
    assert repeated["imported"] == 0
    assert repeated["already_loaded"] == 4
    assert [record["project_id"] for record in first["records"]] == [
        record["project_id"] for record in repeated["records"]
    ]
    assert (await publish(client)).status_code == 409
    assert (await directory(client))["posters"] == []
    admin_rows = (await client.get("/api/v1/projects/?event_id=asm")).json()
    assert len(admin_rows) == 4
    assert all(not row["Is_Published"] for row in admin_rows)
    assert sum(row["Roster_Review_Required"] for row in admin_rows) == 3
    async with sessions() as db:
        changed = prepare_submissions(
            source_fixture(), "source.xlsx", "different snapshot", "Form Responses 1"
        )
        with pytest.raises(ValueError, match="existing roster differs"):
            await stage_submissions(db, "asm", changed)
        await db.rollback()


@pytest.mark.asyncio
async def test_unresolved_review_rows_cannot_leak_if_marked_published(context):
    from app.services.asm_roster import prepare_submissions, stage_submissions

    client, sessions, _ = context
    submissions = prepare_submissions(
        source_fixture(), "source.xlsx", "hash", "Form Responses 1"
    )
    async with sessions() as db:
        await stage_submissions(db, "asm", submissions)
        await db.commit()
        event = await db.get(Event, "asm")
        event.Public_Poster_Lookup_Enabled = True
        rows = await db.scalars(
            select(Project).where(Project.Roster_Review_Required.is_(True))
        )
        for row in rows:
            row.Is_Published = True
        await db.commit()
    assert (await directory(client))["posters"] == []


@pytest.mark.asyncio
async def test_organizer_can_resolve_source_review_without_publishing(context):
    from app.services.asm_roster import prepare_submissions, stage_submissions

    client, sessions, app = context
    submissions = prepare_submissions(
        source_fixture(), "source.xlsx", "hash", "Form Responses 1"
    )
    async with sessions() as db:
        report = await stage_submissions(db, "asm", submissions)
        await db.commit()
    first, duplicate, talk, _ = report["records"]
    first_url = f"/api/v1/projects/{first['project_id']}/approve-roster"
    admin_override = app.dependency_overrides.pop(get_current_admin)
    assert (await client.post(first_url)).status_code == 401
    app.dependency_overrides[get_current_admin] = admin_override
    for record in (first, talk):
        response = await client.post(
            f"/api/v1/projects/{record['project_id']}/approve-roster"
        )
        assert response.status_code == 200, response.text
        assert not response.json()["Roster_Review_Required"]
        assert not response.json()["Is_Published"]
        assert response.json()["Roster_Review"]["approved_by"] == "test-admin"
        assert response.json()["Roster_Review"]["approved_at"]
    duplicate_url = f"/api/v1/projects/{duplicate['project_id']}"
    await client.post(f"{duplicate_url}/withdraw")
    assert (await client.post(f"{duplicate_url}/approve-roster")).status_code == 409
    assert (await directory(client))["posters"] == []
    async with sessions() as db:
        repeated = await stage_submissions(db, "asm", submissions)
        await db.commit()
    assert repeated["records"][0]["review"]["approved_by"] == "test-admin"
    assert not repeated["records"][0]["review_required"]
    assert (await publish(client)).status_code == 200
    assert len((await directory(client))["posters"]) == 3
    assert "approved_by" not in str(await directory(client))


async def scoring_setup(context):
    from app.api.deps import get_current_judge
    from app.models.judge import Judge
    from app.models.rubric import Rubric
    from app.models.rubric_criterion import RubricCriterion

    client, sessions, app = context
    project = await create(
        client,
        Presenter_First_Name="",
        Presenter_Last_Name="",
        Presenter_Display_Name="Jo Chen and Pat Rivera",
    )
    async with sessions() as db:
        judge = Judge(
            Event_ID="asm", First_Name="Test", Last_Name="Judge", Access_Code="ABC123"
        )
        rubric = Rubric(
            Event_ID="asm", Rubric_Name="Poster", Category="Poster", Max_Score=6
        )
        db.add_all([judge, rubric])
        await db.flush()
        criteria = [
            RubricCriterion(Rubric_ID=rubric.Rubric_ID, Criterion_Name=name)
            for name in ("Content", "Communication")
        ]
        db.add_all(criteria)
        await db.commit()
    app.dependency_overrides[get_current_judge] = lambda: judge
    return project, judge, criteria


@pytest.mark.asyncio
async def test_display_names_survive_judging_results_review_and_export(context):
    import io

    from openpyxl import load_workbook

    client, _, _ = context
    project, _, criteria = await scoring_setup(context)
    name = project["Presenter_Display_Name"]
    project_id = project["Project_ID"]
    listed = (await client.get("/api/v1/judge/projects")).json()
    assert listed[0]["Presenter_Display_Name"] == name
    detail = (await client.get(f"/api/v1/judge/projects/{project_id}")).json()
    assert detail["Presenter_Display_Name"] == name
    response = await client.post(
        f"/api/v1/judge/projects/{project_id}/scores",
        json={
            "scores": [
                {"Criterion_ID": c.Criterion_ID, "Score_Value": 3} for c in criteria
            ]
        },
    )
    assert response.status_code == 200, response.text
    results = (await client.get("/api/v1/results/?event_id=asm")).json()
    assert results[0]["Presenter_Display_Name"] == name
    review = (await client.get(f"/api/v1/scores/review/{project_id}")).json()
    assert review["Presenter_Display_Name"] == name
    exported = await client.get("/api/v1/results/export?event_id=asm")
    workbook = load_workbook(io.BytesIO(exported.content))
    assert workbook["Poster Rankings"].cell(2, 4).value == name


@pytest.mark.asyncio
async def test_submission_deactivation_updates_all_counts_and_retains_rows(context):
    from app.models.feedback import Feedback
    from app.models.score import Score

    client, sessions, _ = context
    project, judge, criteria = await scoring_setup(context)
    project_id = project["Project_ID"]
    payload = {
        "scores": [
            {"Criterion_ID": c.Criterion_ID, "Score_Value": 3} for c in criteria
        ],
        "feedback": "Useful work",
    }
    assert (
        await client.post(f"/api/v1/judge/projects/{project_id}/scores", json=payload)
    ).status_code == 200
    summary = await client.get("/api/v1/results/summary?event_id=asm")
    assert summary.status_code == 200, summary.text
    assert summary.json()["Total_Scores_Submitted"] == 1
    assert summary.json()["Scoring_Coverage_Percent"] == 100
    await client.post(f"/api/v1/projects/{project_id}/withdraw")
    summary = (await client.get("/api/v1/results/summary?event_id=asm")).json()
    assert summary["Total_Scores_Submitted"] == 0
    assert summary["Projects_With_Scores"] == 0
    await client.post(f"/api/v1/projects/{project_id}/reinstate")
    response = await client.delete(
        f"/api/v1/scores/submissions/{project_id}/{judge.Judge_ID}"
    )
    assert response.status_code == 204
    assert (await client.get("/api/v1/results/?event_id=asm")).json() == []
    assert (await client.get("/api/v1/results/activity?event_id=asm")).json() == []
    assert (await client.get("/api/v1/judge/my-scores")).json() == []
    assert not (await client.get("/api/v1/judge/projects")).json()[0]["is_scored"]
    summary = (await client.get("/api/v1/results/summary?event_id=asm")).json()
    assert summary["Total_Scores_Submitted"] == 0
    assert summary["Scoring_Coverage_Percent"] == 0
    async with sessions() as db:
        scores = list(await db.scalars(select(Score)))
        assert len(scores) == 2
        assert all(not score.Is_Active for score in scores)
        assert not (await db.scalar(select(Feedback))).Is_Active
        original_ids = {score.Score_ID for score in scores}
    assert (
        await client.post(f"/api/v1/judge/projects/{project_id}/scores", json=payload)
    ).status_code == 200
    async with sessions() as db:
        scores = list(await db.scalars(select(Score)))
        assert {score.Score_ID for score in scores} == original_ids
        assert all(score.Is_Active for score in scores)
    assert (await client.get("/api/v1/results/?event_id=asm")).json()[0][
        "Total_Score"
    ] == 6


@pytest.mark.asyncio
async def test_judge_detail_cannot_read_other_events_or_withdrawn_projects(context):
    client, sessions, _ = context
    project, _, _ = await scoring_setup(context)
    async with sessions() as db:
        db.add(
            Event(
                Event_ID="private-event",
                Event_Name="Private event",
                Event_Date=datetime.date(2026, 3, 13),
                Is_Active=False,
            )
        )
        await db.commit()
    other = await create(client, Event_ID="private-event", Project_Title="Private")
    assert (
        await client.get(f"/api/v1/judge/projects/{other['Project_ID']}")
    ).status_code == 404
    await client.post(f"/api/v1/projects/{project['Project_ID']}/withdraw")
    assert (
        await client.get(f"/api/v1/judge/projects/{project['Project_ID']}")
    ).status_code == 404


@pytest.mark.asyncio
async def test_public_draft_preview_requires_dev_mode_and_preserves_reviews(
    context, monkeypatch
):
    from app.config import settings
    from app.services.asm_roster import prepare_submissions, stage_submissions

    client, sessions, _ = context
    async with sessions() as db:
        await stage_submissions(
            db,
            "asm",
            prepare_submissions(
                source_fixture(), "source.xlsx", "hash", "Form Responses 1"
            ),
        )
        await db.commit()
    assert (await directory(client))["event"] is None
    monkeypatch.setattr(settings, "public_draft_preview", True)
    monkeypatch.setattr(settings, "dev_mode", False)
    assert (await directory(client))["event"] is None
    monkeypatch.setattr(settings, "dev_mode", True)
    payload = await directory(client)
    assert payload["is_draft_preview"]
    assert len(payload["posters"]) == 4
    assert set(payload["posters"][0]) == {
        "Project_Number",
        "Project_Title",
        "Presenter_Display_Name",
        "Institution",
        "Table_Number",
    }
    assert "example.edu" not in str(payload)
    assert "source_row" not in str(payload)
    assert (await publish(client)).status_code == 409
    async with sessions() as db:
        projects = list(await db.scalars(select(Project)))
        assert all(not project.Is_Published for project in projects)
        assert sum(project.Roster_Review_Required for project in projects) == 3
    await create(client, Category="Art", Project_Title="Art stays private")
    project_id = projects[0].Project_ID
    await client.post(f"/api/v1/projects/{project_id}/withdraw")
    payload = await directory(client)
    assert len(payload["posters"]) == 3
    assert "Art stays private" not in str(payload)
    monkeypatch.setattr(settings, "public_draft_preview", False)
    assert (await directory(client))["event"] is None


@pytest.mark.asyncio
async def test_draft_preview_still_requires_one_active_event(context, monkeypatch):
    from app.config import settings

    client, sessions, _ = context
    monkeypatch.setattr(settings, "public_draft_preview", True)
    monkeypatch.setattr(settings, "dev_mode", True)
    await create(client)
    async with sessions() as db:
        db.add(
            Event(
                Event_ID="other",
                Event_Name="Other",
                Event_Date=datetime.date(2026, 10, 31),
                Is_Active=True,
            )
        )
        await db.commit()
    assert (await directory(client))["event"] is None
